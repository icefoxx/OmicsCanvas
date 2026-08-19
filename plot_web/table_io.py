"""Table file I/O helpers: CSV/TSV/TXT + Excel (.xlsx) upload support.

Excel workbooks are converted to CSV (first sheet by default) so existing R
plugins that only read delimited text keep working.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any

TEXT_TABLE_SUFFIXES = {".csv", ".tsv", ".txt"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
# .xls is legacy; support only if a reader is available
LEGACY_EXCEL_SUFFIXES = {".xls"}
TABLE_SUFFIXES = TEXT_TABLE_SUFFIXES | EXCEL_SUFFIXES | LEGACY_EXCEL_SUFFIXES


def is_table_file(path: Path | str) -> bool:
    return Path(path).suffix.lower() in TABLE_SUFFIXES


def is_excel_file(path: Path | str) -> bool:
    return Path(path).suffix.lower() in (EXCEL_SUFFIXES | LEGACY_EXCEL_SUFFIXES)


def _cell_to_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        # avoid 1.0 for integers stored as float
        if value.is_integer():
            return str(int(value))
        return repr(value) if abs(value) >= 1e15 else str(value)
    return str(value)


def _read_xlsx_rows(path: Path, sheet: str | int | None = 0) -> tuple[list[str], list[list[str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Reading .xlsx requires the openpyxl package. "
            "Install with: pip install openpyxl"
        ) from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if isinstance(sheet, str):
            if sheet not in wb.sheetnames:
                raise ValueError(f"Sheet “{sheet}” not found. Available: {wb.sheetnames}")
            ws = wb[sheet]
        else:
            idx = int(sheet or 0)
            if idx < 0 or idx >= len(wb.worksheets):
                raise ValueError(f"Sheet index {idx} out of range ({len(wb.worksheets)} sheets)")
            ws = wb.worksheets[idx]

        matrix: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            if row is None:
                continue
            matrix.append([_cell_to_str(c) for c in row])
    finally:
        wb.close()

    # drop fully empty trailing rows/cols
    while matrix and all(not c.strip() for c in matrix[-1]):
        matrix.pop()
    if not matrix:
        return [], []

    # trim empty trailing columns
    max_cols = 0
    for r in matrix:
        for i in range(len(r) - 1, -1, -1):
            if r[i].strip():
                max_cols = max(max_cols, i + 1)
                break
    matrix = [r[:max_cols] for r in matrix]

    header = matrix[0]
    # uniquify empty / duplicate headers
    columns: list[str] = []
    seen: dict[str, int] = {}
    for i, h in enumerate(header):
        name = (h or "").strip() or f"column_{i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        columns.append(name)

    body = matrix[1:]
    # pad short rows
    rows: list[list[str]] = []
    n = len(columns)
    for r in body:
        if len(r) < n:
            r = r + [""] * (n - len(r))
        rows.append(r[:n])
    return columns, rows


def _read_xls_rows(path: Path, sheet: str | int | None = 0) -> tuple[list[str], list[list[str]]]:
    """Best-effort .xls via pandas if available."""
    try:
        import pandas as pd
    except ImportError as exc:
        raise RuntimeError(
            "Reading legacy .xls requires pandas (+ xlrd). Prefer .xlsx, or convert to CSV."
        ) from exc
    df = pd.read_excel(path, sheet_name=sheet if sheet is not None else 0, dtype=object)
    columns = [str(c) for c in df.columns.tolist()]
    rows: list[list[str]] = []
    for rec in df.fillna("").astype(str).values.tolist():
        rows.append(["" if v == "nan" else v for v in rec])
    return columns, rows


def read_excel_rows(
    path: Path,
    sheet: str | int | None = 0,
) -> tuple[list[str], list[list[str]]]:
    suffix = path.suffix.lower()
    if suffix in EXCEL_SUFFIXES:
        return _read_xlsx_rows(path, sheet=sheet)
    if suffix in LEGACY_EXCEL_SUFFIXES:
        return _read_xls_rows(path, sheet=sheet)
    raise ValueError(f"Not an Excel file: {path.name}")


def write_csv(path: Path, columns: list[str], rows: list[list[str]] | list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(columns)
        for row in rows:
            if isinstance(row, dict):
                writer.writerow([row.get(c, "") for c in columns])
            else:
                writer.writerow(list(row))


def excel_to_csv(
    path: Path,
    *,
    sheet: str | int | None = 0,
    out: Path | None = None,
) -> Path:
    """Convert an Excel workbook sheet to a UTF-8 CSV. Returns CSV path."""
    columns, rows = read_excel_rows(path, sheet=sheet)
    if out is None:
        # e.g. data.xlsx → data.csv (same directory)
        out = path.with_suffix(".csv")
        if out.resolve() == path.resolve():
            out = path.parent / f"{path.stem}_converted.csv"
    write_csv(out, columns, rows)
    return out


def ensure_tabular_csv(
    path: Path,
    *,
    sheet: str | int | None = 0,
    out: Path | None = None,
) -> Path:
    """If path is Excel, convert to CSV and return that path; else return path."""
    if is_excel_file(path):
        return excel_to_csv(path, sheet=sheet, out=out)
    return path


def guess_delimiter(sample: str, filename: str = "") -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in (".tsv", ".txt") and "\t" in sample:
        return "\t"
    if sample.count("\t") > sample.count(","):
        return "\t"
    try:
        dialect = csv.Sniffer().sniff(sample[:4096], delimiters=",\t;")
        return dialect.delimiter
    except Exception:
        return ","


def read_table_as_dicts(
    path: Path,
    *,
    max_rows: int | None = None,
    sheet: str | int | None = 0,
) -> tuple[list[str], list[dict[str, str]], dict[str, Any]]:
    """Read CSV/TSV/TXT or Excel into column names + list of row dicts."""
    suffix = path.suffix.lower()
    meta: dict[str, Any] = {
        "filename": path.name,
        "original_path": str(path),
        "format": suffix.lstrip(".") or "unknown",
    }

    if is_excel_file(path):
        columns, matrix = read_excel_rows(path, sheet=sheet)
        if max_rows is not None:
            matrix = matrix[:max_rows]
        rows = [{columns[i]: row[i] if i < len(row) else "" for i in range(len(columns))} for row in matrix]
        meta.update(
            {
                "delimiter": ",",
                "excel": True,
                "sheet": sheet if sheet is not None else 0,
                "row_count_estimate": len(matrix),
            }
        )
        return columns, rows, meta

    raw = path.read_text(encoding="utf-8", errors="replace")
    sample = raw[:8192]
    delimiter = guess_delimiter(sample, path.name)
    reader = csv.DictReader(raw.splitlines(), delimiter=delimiter)
    columns = list(reader.fieldnames or [])
    rows: list[dict[str, str]] = []
    for idx, row in enumerate(reader):
        if max_rows is not None and idx >= max_rows:
            break
        rows.append({k: (row.get(k) or "") for k in columns})
    n_rows_est = max(0, len(raw.splitlines()) - (1 if columns else 0))
    meta.update(
        {
            "delimiter": "\\t" if delimiter == "\t" else delimiter,
            "excel": False,
            "row_count_estimate": n_rows_est,
        }
    )
    return columns, rows, meta


def sanitize_sheet_name(name: str | None) -> str | int | None:
    if name is None or str(name).strip() == "":
        return 0
    text = str(name).strip()
    if re.fullmatch(r"\d+", text):
        return int(text)
    return text
